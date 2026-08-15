"""Journal-style figure: three-layer learning memory (投稿 §3.6 / 图 2).

Outputs:
  docs/figures/memory_three_layers.pdf
  docs/figures/memory_three_layers.png
  docs/figures/memory_three_layers.svg
  docs/figures/memory_three_layers.pptx
"""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from matplotlib.patches import FancyArrowPatch, Rectangle

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "docs" / "figures"
OUT.mkdir(parents=True, exist_ok=True)
STEM = OUT / "memory_three_layers"

# Full-column: 160 mm x 108 mm
W_MM, H_MM = 160.0, 108.0
NAVY = "#2F4A6D"
EDGE = "#2F4A6D"
FILL_RAIL = "#E8EEF4"
FILL_BOX = "#FFFFFF"
FILL_HEAD = "#2F4A6D"
TEXT = "#1A1A1A"
TEXT_W = "#FFFFFF"

SIMSUN = FontProperties(fname=r"C:\Windows\Fonts\simsun.ttc")
SIMHEI = FontProperties(fname=r"C:\Windows\Fonts\simhei.ttf")


def box(ax, x, y, w, h, facecolor, edgecolor=EDGE, lw=0.5, z=2):
    ax.add_patch(
        Rectangle(
            (x, y),
            w,
            h,
            facecolor=facecolor,
            edgecolor=edgecolor,
            linewidth=lw,
            joinstyle="miter",
            zorder=z,
        )
    )


def label(ax, x, y, text, *, size=8, bold=False, color=TEXT, ha="center", va="center"):
    ax.text(
        x,
        y,
        text,
        fontproperties=SIMHEI if bold else SIMSUN,
        fontsize=size,
        color=color,
        ha=ha,
        va=va,
        zorder=6,
    )


def arrow(ax, x1, y1, x2, y2, *, dashed=False):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1),
            (x2, y2),
            arrowstyle="-|>",
            mutation_scale=7,
            linewidth=0.5,
            color=NAVY,
            linestyle=(0, (2.4, 1.6)) if dashed else "solid",
            shrinkA=0,
            shrinkB=0,
            zorder=4,
        )
    )


def layer_box(ax, x, y, w, h, title, subtitle):
    """Navy header + white body, same idiom as fig1 subsystem columns."""
    head_h = 6.4
    box(ax, x, y, w, h, FILL_BOX)
    box(ax, x, y, w, head_h, FILL_HEAD, edgecolor=NAVY)
    label(ax, x + w / 2, y + head_h / 2, title, size=8, bold=True, color=TEXT_W)
    label(ax, x + w / 2, y + head_h + (h - head_h) / 2, subtitle, size=7.5)


def draw_matplotlib() -> None:
    fig = plt.figure(figsize=(W_MM / 25.4, H_MM / 25.4), dpi=300, facecolor="white")
    ax = fig.add_axes((0, 0, 1, 1))
    ax.set_xlim(0, W_MM)
    ax.set_ylim(H_MM, 0)
    ax.set_axis_off()
    ax.set_facecolor("white")

    # --- geometry ---
    lw, rw = 64.0, 64.0
    lx = 10.0
    rx = 86.0
    bw = 16.4
    turn_w, turn_h = 40.0, 10.0
    turn_x = (W_MM - turn_w) / 2
    turn_y = 4.5

    y1 = 24.5  # L1 / Buf
    y2 = 50.5  # L2 / Job
    y3 = 76.5  # L3 (right)
    y_p = 96.0  # Prompt (full width)
    prompt_h = 8.5

    lcx = lx + lw / 2
    rcx = rx + rw / 2
    tcx = turn_x + turn_w / 2

    # --- nodes ---
    box(ax, turn_x, turn_y, turn_w, turn_h, FILL_RAIL)
    label(ax, tcx, turn_y + turn_h / 2, "每轮对话", size=8, bold=True)
    label(ax, lcx, turn_y + turn_h / 2, "会话内：直接拼入提示词", size=7.5, bold=True)
    label(ax, rcx, turn_y + turn_h / 2, "跨会话：先落库，再检索", size=7.5, bold=True)

    layer_box(ax, lx, y1, lw, bw, "L1 短期记忆", "最近原文，按 token 预算滑动")
    box(ax, rx, y1, rw, bw, FILL_BOX)
    label(ax, rcx, y1 + 6.2, "缓冲队列", size=8, bold=True)
    label(ax, rcx, y1 + 12.2, "攒够轮数或静默超时", size=7.5)

    layer_box(ax, lx, y2, lw, bw, "L2 会话摘要", "滑出窗口的历史，增量结构化压缩")
    box(ax, rx, y2, rw, bw, FILL_BOX)
    label(ax, rcx, y2 + bw / 2, "后台批量抽取", size=8, bold=True)

    layer_box(ax, rx, y3, rw, bw, "L3 长期记忆库", "跨会话事实，按用户与课程隔离")
    box(ax, lx, y_p, rx + rw - lx, prompt_h, FILL_RAIL)
    label(ax, (lx + rx + rw) / 2, y_p + prompt_h / 2, "本轮提示词", size=8, bold=True)

    hub_y = turn_y + turn_h + 3.0
    ax.plot([tcx, tcx], [turn_y + turn_h, hub_y], color=NAVY, lw=0.5, zorder=4)
    ax.plot([lcx, rcx], [hub_y, hub_y], color=NAVY, lw=0.5, zorder=4)
    arrow(ax, lcx, hub_y, lcx, y1)
    arrow(ax, rcx, hub_y, rcx, y1)

    arrow(ax, lcx, y1 + bw, lcx, y2)
    label(ax, lcx + 3.2, (y1 + bw + y2) / 2, "超出窗口的历史", size=7, ha="left")

    arrow(ax, rcx, y1 + bw, rcx, y2)
    arrow(ax, rcx, y2 + bw, rcx, y3)
    label(ax, rcx + 3.2, (y2 + bw + y3) / 2, "写入记忆库", size=7, ha="left")

    # L1 / L2 直接拼入；L3 从记忆库检索后再拼入
    arrow(ax, lcx, y2 + bw, lcx, y_p)
    label(ax, lcx + 3.2, (y2 + bw + y_p) / 2, "直接拼入", size=7, ha="left")
    arrow(ax, rcx, y3 + bw, rcx, y_p)
    label(ax, rx - 2.0, y3 + 9.5, "从记忆库检索", size=7, ha="right")
    label(ax, rx - 2.0, y3 + 13.1, "语义 + 时间衰减 + 矛盾抑制", size=7, ha="right")
    rail_x = lx - 4.0
    mid_p = y_p + prompt_h / 2
    ax.plot(
        [lx, rail_x, rail_x],
        [y1 + bw / 2, y1 + bw / 2, mid_p],
        color=NAVY,
        lw=0.5,
        zorder=4,
        clip_on=False,
    )
    arrow(ax, rail_x, mid_p, lx, mid_p)

    plt.rcParams["pdf.fonttype"] = 42
    plt.rcParams["ps.fonttype"] = 42
    fig.savefig(STEM.with_suffix(".pdf"), facecolor="white")
    fig.savefig(STEM.with_suffix(".png"), dpi=300, facecolor="white")
    fig.savefig(STEM.with_suffix(".svg"), facecolor="white")
    plt.close(fig)


def draw_pptx() -> None:
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Mm, Pt

    prs = Presentation()
    prs.slide_width = Mm(W_MM)
    prs.slide_height = Mm(H_MM)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    navy = RGBColor(0x2F, 0x4A, 0x6D)
    white = RGBColor(0xFF, 0xFF, 0xFF)
    text_c = RGBColor(0x1A, 0x1A, 0x1A)
    fill_rail = RGBColor(0xE8, 0xEE, 0xF4)
    fill_box = RGBColor(0xFF, 0xFF, 0xFF)
    fill_head = RGBColor(0x2F, 0x4A, 0x6D)

    def set_run(run, text, size=8, bold=False, color=text_c, cn="宋体"):
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Times New Roman"
        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = etree.SubElement(rPr, qn("a:ea"))
        ea.set("typeface", "黑体" if bold else cn)

    def add_rect(x, y, w, h, fill, text="", *, size=8, bold=False, color=text_c, lw=0.5):
        shape = slide.shapes.add_shape(1, Mm(x), Mm(y), Mm(w), Mm(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = navy
        shape.line.width = Pt(lw)
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        tf.anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.clear()
        if text:
            run = p.add_run()
            set_run(run, text, size=size, bold=bold, color=color)
        return shape

    def add_tb(x, y, w, h, text, *, size=8, bold=False, color=text_c):
        shape = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        set_run(run, text, size=size, bold=bold, color=color)
        return shape

    def add_line(x1, y1, x2, y2, *, dashed=False):
        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Mm(x1), Mm(y1), Mm(x2), Mm(y2)
        )
        conn.line.color.rgb = navy
        conn.line.width = Pt(0.5)
        if dashed:
            conn.line.prstDash = "dash"
        return conn

    lw, rw = 64.0, 64.0
    lx, rx = 10.0, 86.0
    bw = 16.4
    turn_w, turn_h, turn_y = 40.0, 10.0, 4.5
    turn_x = (W_MM - turn_w) / 2
    y1, y2, y3, y_p = 24.5, 50.5, 76.5, 96.0
    prompt_h = 8.5
    lcx, rcx, tcx = lx + lw / 2, rx + rw / 2, turn_x + turn_w / 2
    hub_y = turn_y + turn_h + 3.0
    head_h = 6.4
    mid_p = y_p + prompt_h / 2
    rail_x = lx - 4.0

    add_rect(turn_x, turn_y, turn_w, turn_h, fill_rail, "每轮对话", bold=True)
    add_tb(lx, y1 - 5.2, lw, 5, "会话内：直接拼入提示词", size=7.5, bold=True)
    add_tb(rx, y1 - 5.2, rw, 5, "跨会话：先落库，再检索", size=7.5, bold=True)

    add_rect(lx, y1, lw, bw, fill_box)
    add_rect(lx, y1, lw, head_h, fill_head, "L1 短期记忆", bold=True, color=white)
    add_tb(lx, y1 + head_h, lw, bw - head_h, "最近原文，按 token 预算滑动", size=7.5)

    add_rect(rx, y1, rw, bw, fill_box)
    add_tb(rx, y1 + 1.5, rw, 6.5, "缓冲队列", bold=True)
    add_tb(rx, y1 + 8.5, rw, 6.5, "攒够轮数或静默超时", size=7.5)

    add_rect(lx, y2, lw, bw, fill_box)
    add_rect(lx, y2, lw, head_h, fill_head, "L2 会话摘要", bold=True, color=white)
    add_tb(lx, y2 + head_h, lw, bw - head_h, "滑出窗口的历史，增量结构化压缩", size=7.5)

    add_rect(rx, y2, rw, bw, fill_box, "后台批量抽取", bold=True)

    add_rect(rx, y3, rw, bw, fill_box)
    add_rect(rx, y3, rw, head_h, fill_head, "L3 长期记忆库", bold=True, color=white)
    add_tb(rx, y3 + head_h, rw, bw - head_h, "跨会话事实，按用户与课程隔离", size=7.5)

    add_rect(lx, y_p, rx + rw - lx, prompt_h, fill_rail, "本轮提示词", bold=True)

    add_line(tcx, turn_y + turn_h, tcx, hub_y)
    add_line(lcx, hub_y, rcx, hub_y)
    add_line(lcx, hub_y, lcx, y1)
    add_line(rcx, hub_y, rcx, y1)
    add_line(lcx, y1 + bw, lcx, y2)
    add_line(rcx, y1 + bw, rcx, y2)
    add_line(rcx, y2 + bw, rcx, y3)
    add_line(lcx, y2 + bw, lcx, y_p)
    add_line(rcx, y3 + bw, rcx, y_p)
    add_line(lx, y1 + bw / 2, rail_x, y1 + bw / 2)
    add_line(rail_x, y1 + bw / 2, rail_x, mid_p)
    add_line(rail_x, mid_p, lx, mid_p)

    add_tb(lcx + 2.0, (y1 + bw + y2) / 2 - 2.2, 28, 5, "超出窗口的历史", size=7)
    add_tb(rcx + 2.0, (y2 + bw + y3) / 2 - 2.2, 22, 5, "写入记忆库", size=7)
    add_tb(lcx + 2.0, (y2 + bw + y_p) / 2 - 2.2, 18, 5, "直接拼入", size=7)
    add_tb(lx + 8, y3 + bw + 2.5, 70, 5, "从记忆库检索", size=7)
    add_tb(lx + 8, y3 + bw + 6.2, 70, 5, "语义 + 时间衰减 + 矛盾抑制", size=7)

    prs.save(STEM.with_suffix(".pptx"))


if __name__ == "__main__":
    draw_matplotlib()
    draw_pptx()
    print(f"wrote {STEM}.pdf / .png / .svg / .pptx")
